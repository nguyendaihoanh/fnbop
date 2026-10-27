# -*- coding: utf-8 -*-
"""
Test Suite: Audit Ledger Engine Triggers, Backup Automation, and Excel Warning
Validates:
1. SQLite Triggers on `so_cai_luong` strictly abort UPDATE and DELETE operations.
2. Retention policy in `cleanup_old_backups` purges files older than 30 days while keeping recent ones.
3. Online atomic `safe_sqlite_backup` preserves WAL consistency.
4. Financial vs Audit separation in `calculate_staff_payslip_breakdown`.
5. Warning banner & staff tagging in `export_payroll_excel` when hanging shifts exist.
"""
import io
import os
import sys
import time
import shutil
import sqlite3
import tempfile
import pytest

# Ensure backend modules can be imported
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from backend.services.db import (
    init_db,
    get_db_connection,
    record_so_cai_luong,
    list_so_cai_luong,
    safe_sqlite_backup,
    cleanup_old_backups,
    upsert_staff_account
)
from backend.services.payroll_calc import calculate_staff_payslip_breakdown
from backend.services.excel_manager import export_payroll_excel, HAS_OPENPYXL
if HAS_OPENPYXL:
    import openpyxl


@pytest.fixture
def temp_db_path():
    """Provides a fresh isolated SQLite database with schema initialized."""
    with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as f:
        path = f.name
    init_db(path)
    yield path
    if os.path.exists(path):
        try:
            os.remove(path)
        except Exception:
            pass
    for ext in ["-wal", "-shm"]:
        if os.path.exists(path + ext):
            try:
                os.remove(path + ext)
            except Exception:
                pass


def test_so_cai_luong_immutability_triggers(temp_db_path):
    """
    Verify that BEFORE UPDATE and BEFORE DELETE triggers on `so_cai_luong`
    raise ABORT and prevent tampering with the audit ledger at the database engine level.
    """
    rec_id = record_so_cai_luong(
        tenant_id="default",
        nhan_vien_id="EMP_001",
        staff_name="Nguyễn Văn A",
        loai_su_kien="THUONG",
        so_tien_thay_doi=100000,
        ly_do="Thưởng chuyên cần tuần",
        nguoi_tao="Manager",
        db_path=temp_db_path
    )
    assert rec_id > 0

    conn = get_db_connection(temp_db_path)

    # 1. Attempt UPDATE -> Must be aborted by SQLite trigger
    with pytest.raises((sqlite3.OperationalError, sqlite3.IntegrityError)) as exc_update:
        conn.execute("UPDATE so_cai_luong SET so_tien_thay_doi = 500000 WHERE id = ?", (rec_id,))
    assert "bat bien" in str(exc_update.value).lower() or "abort" in str(exc_update.value).lower()

    # 2. Attempt DELETE -> Must be aborted by SQLite trigger
    with pytest.raises((sqlite3.OperationalError, sqlite3.IntegrityError)) as exc_delete:
        conn.execute("DELETE FROM so_cai_luong WHERE id = ?", (rec_id,))
    assert "bat bien" in str(exc_delete.value).lower() or "abort" in str(exc_delete.value).lower()

    # 3. Verify record remains unchanged
    cur = conn.cursor()
    cur.execute("SELECT so_tien_thay_doi, ly_do FROM so_cai_luong WHERE id = ?", (rec_id,))
    row = cur.fetchone()
    conn.close()

    assert row is not None
    assert row[0] == 100000
    assert row[1] == "Thưởng chuyên cần tuần"


def test_cleanup_old_backups():
    """
    Verify cleanup_old_backups purges files older than retention_days
    and preserves files within retention window.
    """
    temp_dir = tempfile.mkdtemp(prefix="test_backup_retention_")
    try:
        now = time.time()
        # File 1: 45 days old -> should be deleted
        old_file = os.path.join(temp_dir, "fnbops_backup_20260101_000000.db")
        with open(old_file, "w") as f:
            f.write("old backup content")
        os.utime(old_file, (now - 45 * 86400, now - 45 * 86400))

        # File 2: 10 days old -> should be kept
        recent_file = os.path.join(temp_dir, "fnbops_backup_20260301_000000.db")
        with open(recent_file, "w") as f:
            f.write("recent backup content")
        os.utime(recent_file, (now - 10 * 86400, now - 10 * 86400))

        # File 3: Non-backup file -> should be ignored
        other_file = os.path.join(temp_dir, "notes.txt")
        with open(other_file, "w") as f:
            f.write("some notes")

        deleted = cleanup_old_backups(temp_dir, retention_days=30)

        assert old_file in deleted
        assert not os.path.exists(old_file)
        assert os.path.exists(recent_file)
        assert os.path.exists(other_file)
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


def test_safe_sqlite_backup(temp_db_path):
    """
    Verify safe_sqlite_backup copies database online via native backup API
    and the resulting backup is readable and consistent.
    """
    # Insert test data into source DB
    record_so_cai_luong(
        tenant_id="default",
        nhan_vien_id="EMP_002",
        staff_name="Trần Thị B",
        loai_su_kien="UNG_LUONG",
        so_tien_thay_doi=200000,
        ly_do="Tạm ứng xăng xe",
        db_path=temp_db_path
    )

    with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as bf:
        backup_path = bf.name

    try:
        res_path = safe_sqlite_backup(src_db_path=temp_db_path, backup_db_path=backup_path)
        assert res_path == backup_path
        assert os.path.exists(backup_path)

        # Inspect backup content
        b_conn = sqlite3.connect(backup_path)
        b_cur = b_conn.cursor()
        b_cur.execute("SELECT staff_name, so_tien_thay_doi, loai_su_kien FROM so_cai_luong")
        row = b_cur.fetchone()
        b_conn.close()

        assert row is not None
        assert row[0] == "Trần Thị B"
        assert row[1] == 200000
        assert row[2] == "UNG_LUONG"
    finally:
        if os.path.exists(backup_path):
            os.remove(backup_path)


def test_payroll_calc_audit_vs_financial_separation():
    """
    Verify that DIEU_CHINH_GIO and DOI_MUC_LUONG do not double-count money or hours
    in calculate_staff_payslip_breakdown, while THUONG, PHAT, UNG_LUONG are properly calculated.
    """
    staff_name = "Hoàng Long"
    target_month = "9"

    # Mock so_cai_luong entries
    so_cai_entries = [
        # Financial entries
        {"staff_name": staff_name, "loai_su_kien": "THUONG", "so_tien_thay_doi": 150000, "created_at": "2026-09-10 10:00:00"},
        {"staff_name": staff_name, "loai_su_kien": "PHAT", "so_tien_thay_doi": 50000, "created_at": "2026-09-12 11:00:00"},
        {"staff_name": staff_name, "loai_su_kien": "UNG_LUONG", "so_tien_thay_doi": 200000, "created_at": "2026-09-15 15:00:00"},
        # Audit log only entries (MUST NOT alter money)
        {"staff_name": staff_name, "loai_su_kien": "DIEU_CHINH_GIO", "so_gio_thay_doi": 5.0, "so_tien_thay_doi": 85000, "created_at": "2026-09-16 16:00:00"},
        {"staff_name": staff_name, "loai_su_kien": "DOI_MUC_LUONG", "so_gio_thay_doi": 0.0, "so_tien_thay_doi": 25000, "created_at": "2026-09-20 09:00:00"},
    ]

    breakdown = calculate_staff_payslip_breakdown(
        staff_name=staff_name,
        month=target_month,
        data_json={},
        so_cai_luong=so_cai_entries
    )

    # THUONG should be in eval_bonus
    assert breakdown["eval_bonus"] == 150000
    # PHAT should be in eval_deductions
    assert breakdown["eval_deductions"] == 50000
    # UNG_LUONG should be in advance_amount
    assert breakdown["advance_amount"] == 200000

    # DIEU_CHINH_GIO / DOI_MUC_LUONG must NOT inflate standard_pay or bonus
    assert breakdown["standard_pay"] == 0
    assert breakdown["total_hours"] == 0.0


def test_excel_export_warning_banner_and_dynamic_rows(temp_db_path):
    """
    Verify export_payroll_excel renders:
    1. Warning banner at row 3 when hanging shifts exist.
    2. Dynamic header at row 4, first data row at row 5.
    3. Staff name tagged with ⚠️(Treo).
    4. Total row SUM formulas referencing row 5.
    """
    if not HAS_OPENPYXL:
        pytest.skip("openpyxl is not installed")

    # Seed staff account in SQLite
    upsert_staff_account(
        username="long_bar",
        employee_id="EMP_003",
        full_name="Hoàng Long",
        role="STAFF",
        position="BAR",
        hourly_rate=17000,
        db_path=temp_db_path
    )

    # Insert a pending hanging shift into attendance_logs in SQLite
    conn = get_db_connection(temp_db_path)
    conn.execute("""
        INSERT INTO attendance_logs (
            tenant_id, username, full_name, date, shift, checkin, is_hanging_pending, status
        ) VALUES ('default', 'long_bar', 'Hoàng Long', '15/09/2026', 'Ca Tối', '2026-09-15 17:00:00', 1, 'QUEN_CHECKOUT')
    """)
    conn.commit()
    conn.close()

    excel_bytes = export_payroll_excel(month_str="9", db_path=temp_db_path)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active

    # Check row 3 is the warning banner
    row3_val = str(ws["A3"].value or "")
    assert "CẢNH BÁO QUYẾT TOÁN" in row3_val
    assert "QUEN_CHECKOUT" in row3_val

    # Check row 4 is the column headers
    assert ws["A4"].value == "STT"
    assert ws["C4"].value == "Tên NV"

    # Check row 5 is the data row for Hoàng Long
    assert ws["B5"].value == "long_bar"
    c_name_val = str(ws["C5"].value or "")
    assert "⚠️(Treo)" in c_name_val

    # Check total row formulas reference row 5
    # Total row should be at row 6
    assert ws["A6"].value == "TỔNG CỘNG"
    assert ws["E6"].value == "=SUM(E5:E5)"
    assert ws["F6"].value == "=SUM(F5:F5)"
    assert ws["M6"].value == "=SUM(M5:M5)"
