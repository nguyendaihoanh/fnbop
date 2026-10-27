"""
Unit & Integration Tests for Workstream 1: SQLite Schema & Excel Manager
Tests SQLite WAL DDL, JSON-to-SQLite Migration, openpyxl Payroll Export/Import,
Staff Template Generation, and FastAPI Excel Router.
"""

import io
import os
import sys
import tempfile
import sqlite3
import pytest
import openpyxl
from fastapi.testclient import TestClient

# Ensure backend and project root are in sys.path
_current_dir = os.path.dirname(os.path.abspath(__file__))
_project_root = os.path.dirname(_current_dir)
_backend_dir = os.path.join(_project_root, "backend")
if _backend_dir not in sys.path:
    sys.path.insert(0, _backend_dir)
if _project_root not in sys.path:
    sys.path.insert(0, _project_root)

from services.db import (
    init_db,
    get_db_connection,
    upsert_staff_account,
    get_staff_by_username,
    list_staff_accounts,
    insert_evaluation,
    list_evaluations,
    insert_daily_closing,
    get_daily_closings,
    list_daily_closings,
    get_daily_closing_by_id,
    migrate_seed_data_from_json,
    DEFAULT_TENANT_ID
)
from services.excel_manager import (
    export_payroll_excel,
    generate_staff_template,
    import_staff_from_excel
)
from main import app


@pytest.fixture
def temp_db():
    """Provides an isolated SQLite database file in WAL mode."""
    with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as f:
        db_path = f.name

    init_db(db_path)
    yield db_path

    for ext in ["", "-wal", "-shm"]:
        p = f"{db_path}{ext}"
        if os.path.exists(p):
            try:
                os.remove(p)
            except Exception:
                pass


from routers.fnb_ops import create_access_token


@pytest.fixture
def client():
    """FastAPI TestClient fixture with Admin authorization."""
    c = TestClient(app)
    token = create_access_token({"sub": "admin", "role": "ADMIN", "name": "Chủ Quán"})
    c.headers["Authorization"] = f"Bearer {token}"
    return c


# ==============================================================================
# 1. SQLITE SCHEMA & DDL VALIDATION
# ==============================================================================

def test_sqlite_schema_tables_and_columns(temp_db):
    """
    Verifies that staff_accounts, daily_closings, and evaluations tables
    contain all required columns per Master Plan DDL specification.
    """
    conn = get_db_connection(temp_db)
    cur = conn.cursor()

    # 1. staff_accounts
    cur.execute("PRAGMA table_info(staff_accounts);")
    staff_cols = {r[1]: r[2].upper() for r in cur.fetchall()}
    required_staff = [
        "username", "display_name", "role", "pin_hash",
        "email", "phone", "hourly_rate", "is_active", "created_at"
    ]
    for col in required_staff:
        assert col in staff_cols, f"Column '{col}' missing from staff_accounts schema!"

    # 2. daily_closings
    cur.execute("PRAGMA table_info(daily_closings);")
    closing_cols = {r[1]: r[2].upper() for r in cur.fetchall()}
    required_closing = [
        "id", "report_date", "closing_time", "bill_pos", "cash_amount",
        "transfer_amount", "expense_amount", "expense_detail", "total_real",
        "diff_amount", "diff_status", "sender_name", "approver_name", "notes", "created_at"
    ]
    for col in required_closing:
        assert col in closing_cols, f"Column '{col}' missing from daily_closings schema!"

    # 3. evaluations (Must have INTEGER PRIMARY KEY AUTOINCREMENT)
    cur.execute("PRAGMA table_info(evaluations);")
    eval_cols = {r[1]: r[2].upper() for r in cur.fetchall()}
    required_eval = [
        "id", "staff_name", "type", "amount", "reason",
        "date", "month", "manager", "status", "created_at"
    ]
    for col in required_eval:
        assert col in eval_cols, f"Column '{col}' missing from evaluations schema!"

    # Verify id is integer
    assert "INT" in eval_cols["id"], f"evaluations.id type must be INTEGER, got {eval_cols['id']}"

    conn.close()


# ==============================================================================
# 2. DATA MIGRATION FROM JSON TO SQLITE
# ==============================================================================

def test_migrate_seed_data_from_json(temp_db):
    """
    Tests reading staff_auth.json (standard 4 staff) and data/data.json
    and upserting them cleanly into SQLite. Verifies idempotency.
    """
    res = migrate_seed_data_from_json(db_path=temp_db, force=True)
    assert res["success"] is True
    assert res["migrated_staff"] >= 4, f"Expected at least 4 staff migrated, got {res['migrated_staff']}"

    # Verify key staff records
    stf_mgr = get_staff_by_username("manager", db_path=temp_db)
    assert stf_mgr is not None
    assert stf_mgr["display_name"] == "Quản Lý Ca" or stf_mgr["full_name"] == "Quản Lý Ca"
    assert stf_mgr["role"] == "MANAGER"
    assert stf_mgr["is_active"] is True
    assert stf_mgr["pin_hash"].startswith("pbkdf2_sha256$")

    stf_pha = get_staff_by_username("phache", db_path=temp_db)
    assert stf_pha is not None
    assert stf_pha["role"] == "STAFF"

    stf_nv = get_staff_by_username("nhanvien", db_path=temp_db)
    assert stf_nv is not None
    assert stf_nv["role"] == "STAFF"

    # Idempotency check: running again without force should skip
    res2 = migrate_seed_data_from_json(db_path=temp_db, force=False)
    assert res2["skipped"] is True


# ==============================================================================
# 3. STAFF ACCOUNTS CRUD & REPOSITORY
# ==============================================================================

def test_staff_accounts_upsert_and_list(temp_db):
    """
    Tests inserting, updating, and listing staff accounts with extended fields.
    """
    upsert_staff_account(
        username="barista_linh",
        full_name="Nguyễn Thùy Linh",
        display_name="Thùy Linh Barista",
        pin_hash="test_pin_hash",
        role="STAFF",
        position="BARISTA",
        email="linh@wfm.vn",
        phone="0912345678",
        hourly_rate=19000,
        is_active=True,
        db_path=temp_db
    )

    linh = get_staff_by_username("barista_linh", db_path=temp_db)
    assert linh is not None
    assert linh["display_name"] == "Thùy Linh Barista"
    assert linh["hourly_rate"] == 19000
    assert linh["is_active"] is True
    assert linh["email"] == "linh@wfm.vn"

    # Update hourly rate & status
    upsert_staff_account(
        username="barista_linh",
        full_name="Nguyễn Thùy Linh",
        display_name="Thùy Linh Barista",
        role="STAFF",
        hourly_rate=21000,
        is_active=False,
        db_path=temp_db
    )

    linh_updated = get_staff_by_username("barista_linh", db_path=temp_db)
    assert linh_updated["hourly_rate"] == 21000
    assert linh_updated["is_active"] is False


# ==============================================================================
# 4. DAILY CLOSINGS CRUD
# ==============================================================================

def test_daily_closings_crud(temp_db):
    """
    Tests inserting and querying POS cash closing records.
    """
    closing_data = {
        "report_date": "2026-09-24",
        "closing_time": "23:15:00",
        "bill_pos": 3500000.0,
        "cash_amount": 1500000.0,
        "transfer_amount": 1950000.0,
        "expense_amount": 50000.0,
        "expense_detail": "Mua đá lạnh và chanh tươi",
        "total_real": 3500000.0,
        "diff": 0.0,
        "diff_amount": 0.0,
        "diff_status": "KHỚP",
        "is_warning": False,
        "sender_name": "nhanvien",
        "approver_name": "admin",
        "notes": "Kiểm két ca tối khớp 100%"
    }

    rec = insert_daily_closing(closing_data, db_path=temp_db)
    assert rec["id"] is not None
    closing_id = rec["id"]

    # Query list
    closings = list_daily_closings(limit=10, db_path=temp_db)
    assert len(closings) == 1
    assert closings[0]["diff_status"] == "KHỚP"
    assert closings[0]["total_real"] == 3500000.0

    # Query single by ID
    single = get_daily_closing_by_id(closing_id, db_path=temp_db)
    assert single is not None
    assert single["notes"] == "Kiểm két ca tối khớp 100%"


# ==============================================================================
# 5. EVALUATIONS AUTOINCREMENT & QUERIES
# ==============================================================================

def test_evaluations_autoincrement(temp_db):
    """
    Tests autoincrement ID generation for evaluations.
    """
    id1 = insert_evaluation({
        "staff_name": "Nhân Viên",
        "date": "2026-09-24",
        "month": "9",
        "type": "bonus",
        "amount": 100000,
        "reason": "Phục vụ khách nhiệt tình",
        "manager": "Quản Lý Ca"
    }, db_path=temp_db)

    id2 = insert_evaluation({
        "staff_name": "Pha Chế",
        "date": "2026-09-24",
        "month": "9",
        "type": "penalty",
        "amount": 50000,
        "reason": "Đi trễ ca tối",
        "manager": "Quản Lý Ca"
    }, db_path=temp_db)

    assert isinstance(id1, int)
    assert isinstance(id2, int)
    assert id2 > id1

    evals = list_evaluations(month="9", db_path=temp_db)
    assert len(evals) == 2


# ==============================================================================
# 6. EXPORT PAYROLL EXCEL
# ==============================================================================

def test_export_payroll_excel(temp_db):
    """
    Verifies export_payroll_excel:
    - Returns valid .xlsx file bytes
    - Brand header color: Slate Navy #18222F with white text
    - Exact columns: STT, Mã NV, Tên NV, Chức vụ, Tổng ca, Tổng giờ,
      Lương chuẩn, Tiền OT, Phụ cấp ngày lễ, Thưởng, Khấu trừ/Phạt, Ứng lương, Thực lĩnh
    - Net Pay Excel Formula: =G{row}+H{row}+I{row}+J{row}-K{row}-L{row}
    - Currency number format: #,##0 "đ"
    - Total row with SUM formulas
    """
    migrate_seed_data_from_json(db_path=temp_db, force=True)

    excel_bytes = export_payroll_excel(month_str="8", db_path=temp_db)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 2000

    # Load into openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=False)
    ws = wb.active
    assert "8" in ws.title

    # Header Row 3 check
    expected_headers = [
        "STT", "Mã NV", "Tên NV", "Chức vụ", "Tổng ca", "Tổng giờ",
        "Lương chuẩn", "Tiền OT", "Phụ cấp ngày lễ", "Thưởng", "Khấu trừ/Phạt", "Ứng lương", "Thực lĩnh"
    ]
    for col_idx, expected in enumerate(expected_headers, 1):
        cell_val = ws.cell(row=3, column=col_idx).value
        assert cell_val == expected, f"Col {col_idx} expected '{expected}', got '{cell_val}'"

        # Check Navy Header fill
        cell_fill = ws.cell(row=3, column=col_idx).fill.start_color.rgb
        assert "18222F" in str(cell_fill).upper(), f"Header cell fill should be 18222F, got {cell_fill}"

    # Check Data Row 4 formula and formatting
    assert ws.max_row >= 5
    first_data_row = 4
    net_formula = ws.cell(row=first_data_row, column=13).value
    assert net_formula == f"=G{first_data_row}+H{first_data_row}+I{first_data_row}+J{first_data_row}-K{first_data_row}-L{first_data_row}"

    std_format = ws.cell(row=first_data_row, column=7).number_format
    assert 'đ' in std_format or '0' in std_format

    # Check Total Row at bottom
    last_row = ws.max_row
    assert ws.cell(row=last_row, column=1).value == "TỔNG CỘNG"
    assert ws.cell(row=last_row, column=13).value == f"=SUM(M4:M{last_row-1})"


# ==============================================================================
# 7. GENERATE STAFF TEMPLATE
# ==============================================================================

def test_generate_staff_template():
    """
    Verifies staff template generation:
    - Returns valid .xlsx bytes
    - Has guide row and required columns
    - Has 3 sample data rows
    """
    template_bytes = generate_staff_template()
    assert isinstance(template_bytes, bytes)
    assert len(template_bytes) > 2000

    wb = openpyxl.load_workbook(io.BytesIO(template_bytes), data_only=True)
    ws = wb.active
    assert ws.title == "DanhSachNhanSu"

    # Check instructions on Row 2
    assert "HƯỚNG DẪN" in str(ws["A2"].value)

    # Check Headers on Row 3
    headers = [ws.cell(row=3, column=c).value for c in range(1, 10)]
    assert "Mã NV" in headers[0]
    assert "Họ và tên" in headers[1]
    assert "Chức vụ" in headers[2]
    assert "Mã PIN" in headers[3]

    # Check 3 sample rows
    assert ws.cell(row=4, column=1).value == "nhanvien01"
    assert ws.cell(row=5, column=1).value == "nhanvien02"
    assert ws.cell(row=6, column=1).value == "nhanvien03"


# ==============================================================================
# 8. IMPORT STAFF FROM EXCEL
# ==============================================================================

def test_import_staff_from_excel(temp_db):
    """
    Tests parsing Excel file bytes and importing staff into SQLite staff_accounts.
    """
    template_bytes = generate_staff_template()
    res = import_staff_from_excel(template_bytes, db_path=temp_db)

    assert res["success"] is True
    assert res["count"] == 3
    assert res["errors"] == []

    # Check accounts created in DB
    p = get_staff_by_username("nhanvien01", db_path=temp_db)
    assert p is not None
    assert p["display_name"] == "Nhân Viên 01"
    assert p["role"] == "MANAGER"

    ny = get_staff_by_username("nhanvien02", db_path=temp_db)
    assert ny is not None
    assert ny["display_name"] == "Nhân Viên 02"

    th = get_staff_by_username("nhanvien03", db_path=temp_db)
    assert th is not None
    assert th["display_name"] == "Nhân Viên 03"


# ==============================================================================
# 9. FASTAPI EXCEL ROUTER ENDPOINTS
# ==============================================================================

def test_excel_api_export_payroll(client, temp_db):
    """
    Tests GET /api/v1/cc/excel/export-payroll endpoint.
    """
    orig_db = os.environ.get("DATABASE_PATH")
    os.environ["DATABASE_PATH"] = temp_db
    try:
        migrate_seed_data_from_json(db_path=temp_db, force=True)
        response = client.get("/api/v1/cc/excel/export-payroll?month=8")
        assert response.status_code == 200
        assert "application/vnd.openxmlformats" in response.headers["content-type"]
        assert "BangLuong_FnBOps_Thang_8.xlsx" in response.headers.get("content-disposition", "")
        assert len(response.content) > 2000

        # Validate file can be loaded
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        assert "Bảng Lương T8" in wb.sheetnames
    finally:
        if orig_db is not None:
            os.environ["DATABASE_PATH"] = orig_db
        else:
            os.environ.pop("DATABASE_PATH", None)


def test_excel_api_staff_template(client):
    """
    Tests GET /api/v1/cc/excel/staff-template endpoint.
    """
    response = client.get("/api/v1/cc/excel/staff-template")
    assert response.status_code == 200
    assert "application/vnd.openxmlformats" in response.headers["content-type"]
    assert "Mau_Nhan_Su_FnBOps.xlsx" in response.headers.get("content-disposition", "")
    assert len(response.content) > 2000

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "DanhSachNhanSu" in wb.sheetnames


def test_excel_api_import_staff(client, temp_db):
    """
    Tests POST /api/v1/cc/excel/import-staff endpoint with valid file and invalid extension.
    """
    orig_db = os.environ.get("DATABASE_PATH")
    orig_auth = os.environ.get("AUTH_FILE")
    with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as f:
        temp_auth = f.name
    import shutil
    real_auth = os.path.join(_project_root, "data", "staff_auth.json")
    if os.path.exists(real_auth):
        shutil.copyfile(real_auth, temp_auth)

    os.environ["DATABASE_PATH"] = temp_db
    os.environ["AUTH_FILE"] = temp_auth
    try:
        template_bytes = generate_staff_template()

        # 1. Valid Excel upload
        files = {
            "file": ("staff_import.xlsx", template_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        }
        res = client.post("/api/v1/cc/excel/import-staff", files=files)
        assert res.status_code == 200
        data = res.json()
        assert data["success"] is True
        assert data["count"] == 3

        # 2. Invalid file extension upload
        bad_files = {
            "file": ("invalid.txt", b"plain text", "text/plain")
        }
        bad_res = client.post("/api/v1/cc/excel/import-staff", files=bad_files)
        assert bad_res.status_code == 400
    finally:
        if orig_db is not None:
            os.environ["DATABASE_PATH"] = orig_db
        else:
            os.environ.pop("DATABASE_PATH", None)
        if orig_auth is not None:
            os.environ["AUTH_FILE"] = orig_auth
        else:
            os.environ.pop("AUTH_FILE", None)
        if os.path.exists(temp_auth):
            try:
                os.remove(temp_auth)
            except Exception:
                pass
